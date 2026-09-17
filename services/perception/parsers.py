"""文件正文解析器（v0.7）：只做“只读提取”，绝不执行、绝不解压。

解析器契约：
- 输入是**已经落到安全临时目录**的本机路径 + 已经过 file_types 校验的类型；
- 输出 ParseResult(text=..., note=..., ok=...)；
- **绝不** eval / exec / os.system / shell / subprocess 执行任何用户文件；
- 只读文本类代码文件（.py / .cpp / .sh / .js ...）永远只被当作字符串读取；
- 可选解析库缺失时返回结构化 note，而不是让整个聊天失败；
- 每个解析器都有自己的页数 / 段落 / 行数 / 列数上限，超限只截断并说明。

同步 CPU/IO 解析（PyMuPDF / openpyxl / python-docx / python-pptx）在
services/file_reader.py 里统一用 asyncio.to_thread() 调用，绝不阻塞事件循环。
"""

from dataclasses import dataclass

from nonebot import logger

from services.perception import file_types
from services.perception.limits import DOCX_MAX_PARAGRAPHS
from services.perception.limits import PDF_MAX_PAGES
from services.perception.limits import PPTX_MAX_SLIDES
from services.perception.limits import XLSX_MAX_COLS
from services.perception.limits import XLSX_MAX_ROWS_PER_SHEET
from services.perception.limits import XLSX_MAX_SHEETS

# 文本类文件的读取编码候选（按顺序尝试；全部失败用 errors="replace" 兜底）
TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1")

# 稳定的失败说明（对模型可见，只描述程序事实）
PDF_NO_TEXT_NOTE = "该 PDF 没有检测到可提取的文本内容。"
PARSE_FAILED_NOTE = "文件内容解析失败。"
EMPTY_TEXT_NOTE = "文件里没有可读取的文本内容。"
BINARY_NOTE = "暂不支持读取该文件类型。"
TRUNCATED_SUFFIX = "\n…（后续内容因长度限制已截断）"


@dataclass(frozen=True)
class ParseResult:
    """一次正文提取的结果（text 已按字符上限截断）。"""

    ok: bool
    text: str = ""
    note: str = ""
    pages: int = 0          # PDF 总页数（其它类型为 0）
    truncated: bool = False


def _clip(text: str, max_chars: int) -> tuple[str, bool]:
    """按字符数截断；返回 (文本, 是否发生截断)。"""
    text = text or ""
    if max_chars <= 0:
        return "", bool(text)
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars] + TRUNCATED_SUFFIX, True


# ===== 纯文本 / 代码 =====


def parse_text(path: str, max_chars: int) -> ParseResult:
    """只读文本内容。代码文件（.py/.cpp/.sh/...）同样只当作字符串读取。"""
    for encoding in TEXT_ENCODINGS:
        try:
            with open(path, "r", encoding=encoding, errors="strict") as handle:
                raw = handle.read(max_chars + 1)
        except UnicodeDecodeError:
            continue
        except OSError as exc:
            logger.warning("[FILE] 文本读取失败（category=io_error）：{}", type(exc).__name__)
            return ParseResult(ok=False, note=PARSE_FAILED_NOTE)
        text, truncated = _clip(raw, max_chars)
        if not text.strip():
            return ParseResult(ok=True, text="", note=EMPTY_TEXT_NOTE)
        return ParseResult(ok=True, text=text, truncated=truncated)

    # 所有严格编码都失败：用替换字符兜底，绝不把文件读失败升级成聊天失败。
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            raw = handle.read(max_chars + 1)
    except OSError as exc:
        logger.warning("[FILE] 文本读取失败（category=io_error）：{}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)
    text, truncated = _clip(raw, max_chars)
    if not text.strip():
        return ParseResult(ok=True, text="", note=EMPTY_TEXT_NOTE)
    return ParseResult(ok=True, text=text, truncated=truncated)


# ===== PDF（PyMuPDF，仅文字层，不做 OCR）=====


def parse_pdf(path: str, max_chars: int) -> ParseResult:
    """提取 PDF 文字层（PyMuPDF）。**不做 OCR**，扫描版返回明确说明。"""
    try:
        # PyMuPDF >= 1.24 推荐 `import pymupdf`；旧版本只有 `fitz` 别名。
        try:
            import pymupdf as fitz
        except ImportError:  # pragma: no cover - 取决于安装的 PyMuPDF 版本
            import fitz  # type: ignore[no-redef]
    except Exception:
        logger.warning("[FILE] 未安装 PyMuPDF，无法解析 PDF")
        return ParseResult(ok=False, note="当前环境未安装 PDF 解析库（PyMuPDF）。")

    try:
        document = fitz.open(path)
    except Exception as exc:
        logger.info("[FILE] PDF 打开失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    try:
        page_total = document.page_count
        limit = min(page_total, PDF_MAX_PAGES)
        chunks: list[str] = []
        used = 0
        truncated = False
        for index in range(limit):
            page = document.load_page(index)
            page_text = page.get_text() or ""
            if not page_text.strip():
                continue
            chunks.append(page_text)
            used += len(page_text)
            if used >= max_chars:
                truncated = True
                break
        if page_total > limit:
            truncated = True
    except Exception as exc:
        logger.info("[FILE] PDF 解析失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)
    finally:
        try:
            document.close()
        except Exception:
            pass

    text = "\n".join(chunks).strip()
    if not text:
        return ParseResult(ok=True, text="", note=PDF_NO_TEXT_NOTE, pages=page_total)
    clipped, clip_truncated = _clip(text, max_chars)
    if page_total > PDF_MAX_PAGES:
        clipped += f"\n（该 PDF 共 {page_total} 页，仅读取前 {PDF_MAX_PAGES} 页）"
    return ParseResult(
        ok=True,
        text=clipped,
        pages=page_total,
        truncated=truncated or clip_truncated,
    )


# ===== DOCX（python-docx：段落 + 表格单元格）=====


def parse_docx(path: str, max_chars: int) -> ParseResult:
    """提取 DOCX 段落与表格单元格文本（保持基本顺序）。"""
    try:
        import docx  # python-docx
    except Exception:
        logger.warning("[FILE] 未安装 python-docx，无法解析 DOCX")
        return ParseResult(ok=False, note="当前环境未安装 DOCX 解析库（python-docx）。")

    try:
        document = docx.Document(path)
    except Exception as exc:
        logger.info("[FILE] DOCX 打开失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    lines: list[str] = []
    paragraphs = 0
    truncated = False
    try:
        for paragraph in document.paragraphs:
            if paragraphs >= DOCX_MAX_PARAGRAPHS:
                truncated = True
                break
            paragraphs += 1
            text = (paragraph.text or "").strip()
            if text:
                lines.append(text)
        if not truncated:
            for table in document.tables:
                for row in table.rows:
                    if paragraphs >= DOCX_MAX_PARAGRAPHS:
                        truncated = True
                        break
                    paragraphs += 1
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
                if truncated:
                    break
    except Exception as exc:
        logger.info("[FILE] DOCX 解析失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    text = "\n".join(lines).strip()
    if not text:
        return ParseResult(ok=True, text="", note=EMPTY_TEXT_NOTE)
    clipped, clip_truncated = _clip(text, max_chars)
    if truncated:
        clipped += f"\n（文档段落数超过 {DOCX_MAX_PARAGRAPHS} 条，仅读取前 {DOCX_MAX_PARAGRAPHS} 条）"
    return ParseResult(ok=True, text=clipped, truncated=truncated or clip_truncated)


# ===== XLSX（openpyxl：read_only，不执行公式、不运行宏）=====


def parse_xlsx(path: str, max_chars: int) -> ParseResult:
    """只读单元格内容（read_only=True）；公式只保留原文，**不计算**、不运行宏。"""
    try:
        import openpyxl
    except Exception:
        logger.warning("[FILE] 未安装 openpyxl，无法解析 XLSX")
        return ParseResult(ok=False, note="当前环境未安装 XLSX 解析库（openpyxl）。")

    try:
        # data_only=False：保留公式原文，绝不触发任何宏或外部链接求值。
        workbook = openpyxl.load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        logger.info("[FILE] XLSX 打开失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    sections: list[str] = []
    truncated = False
    sheet_total = 0
    try:
        sheet_names = list(workbook.sheetnames)
        sheet_total = len(sheet_names)
        for sheet_name in sheet_names[:XLSX_MAX_SHEETS]:
            sheet = workbook[sheet_name]
            lines = [f"Sheet: {sheet_name}"]
            rows = 0
            for row in sheet.iter_rows(
                max_row=XLSX_MAX_ROWS_PER_SHEET, max_col=XLSX_MAX_COLS
            ):
                rows += 1
                cells = []
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    cells.append(f"{cell.coordinate}: {value}")
                if cells:
                    lines.append(" | ".join(cells))
            if rows >= XLSX_MAX_ROWS_PER_SHEET:
                truncated = True
                lines.append(f"（该表超过 {XLSX_MAX_ROWS_PER_SHEET} 行，仅读取前 {XLSX_MAX_ROWS_PER_SHEET} 行）")
            sections.append("\n".join(lines))
        if sheet_total > XLSX_MAX_SHEETS:
            truncated = True
            sections.append(f"（该文件共 {sheet_total} 个工作表，仅读取前 {XLSX_MAX_SHEETS} 个）")
    except Exception as exc:
        logger.info("[FILE] XLSX 解析失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    text = "\n\n".join(section for section in sections if section.strip()).strip()
    if not text:
        return ParseResult(ok=True, text="", note=EMPTY_TEXT_NOTE)
    clipped, clip_truncated = _clip(text, max_chars)
    return ParseResult(ok=True, text=clipped, truncated=truncated or clip_truncated)


# ===== PPTX（python-pptx：按 slide 提取标题 + 正文）=====


def parse_pptx(path: str, max_chars: int) -> ParseResult:
    """按 slide 顺序提取文本（标题与正文合并为一段）。"""
    try:
        from pptx import Presentation
    except Exception:
        logger.warning("[FILE] 未安装 python-pptx，无法解析 PPTX")
        return ParseResult(ok=False, note="当前环境未安装 PPTX 解析库（python-pptx）。")

    try:
        presentation = Presentation(path)
    except Exception as exc:
        logger.info("[FILE] PPTX 打开失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    sections: list[str] = []
    slide_total = 0
    truncated = False
    try:
        slides = list(presentation.slides)
        slide_total = len(slides)
        for number, slide in enumerate(slides[:PPTX_MAX_SLIDES], start=1):
            lines = [f"Slide {number}"]
            for shape in slide.shapes:
                text = getattr(shape, "text", "") or ""
                text = text.strip()
                if text:
                    lines.append(text)
            sections.append("\n".join(lines))
        if slide_total > PPTX_MAX_SLIDES:
            truncated = True
            sections.append(f"（该文件共 {slide_total} 页，仅读取前 {PPTX_MAX_SLIDES} 页）")
    except Exception as exc:
        logger.info("[FILE] PPTX 解析失败（category=malformed）: {}", type(exc).__name__)
        return ParseResult(ok=False, note=PARSE_FAILED_NOTE)

    text = "\n\n".join(section for section in sections if section.strip()).strip()
    if not text:
        return ParseResult(ok=True, text="", note=EMPTY_TEXT_NOTE)
    clipped, clip_truncated = _clip(text, max_chars)
    return ParseResult(ok=True, text=clipped, truncated=truncated or clip_truncated)


# ===== 统一分发 =====

_PARSERS = {
    file_types.PARSER_TEXT: parse_text,
    file_types.PARSER_PDF: parse_pdf,
    file_types.PARSER_DOCX: parse_docx,
    file_types.PARSER_XLSX: parse_xlsx,
    file_types.PARSER_PPTX: parse_pptx,
}


def parse_by_type(parser: str, path: str, max_chars: int) -> ParseResult:
    """按 parser_type 分发到具体解析器；未知类型返回 unsupported 说明。"""
    handler = _PARSERS.get(parser)
    if handler is None:
        return ParseResult(ok=False, note=BINARY_NOTE)
    return handler(path, max_chars)
